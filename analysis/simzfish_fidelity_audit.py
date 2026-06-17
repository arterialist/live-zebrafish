"""Audit current MuJoCo video pipeline against public Z-Robot/simZFish resources.

This is not a behavior run.  It is an implementation-fidelity layer for the
two-paper zebrafish study: it extracts evidence from the cached public
simZFish C/Webots controller, the published Data_Liu_simZFish_2025 resources,
and the current Python/MuJoCo video branch, then writes a scorecard that is
explicit about what is exact, approximate, or still missing.
"""

from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy import io as scipy_io


ROOT = Path(__file__).resolve().parents[1]
ACTIVE_INFERENCE = ROOT.parent / "active-inference"
ZROBOT = ROOT / "analysis" / "cache" / "z_robot"
SIMZFISH_RAW = ZROBOT / "simzfish_raw" / "controllers" / "ZebrafishRobot_mini"
SIMZFISH_DATA = ZROBOT / "simzfish_data" / "Data_Liu_simZFish_2025"
CURRENT_OMR = ACTIVE_INFERENCE / "simulations" / "zebrafish" / "simzfish_omr.py"
CURRENT_ACTION = ACTIVE_INFERENCE / "simulations" / "zebrafish" / "action_latent.py"
CURRENT_CONFIG = ACTIVE_INFERENCE / "simulations" / "zebrafish" / "config.py"
CURRENT_VIDEO = ROOT / "lab" / "video_pipeline.py"
OUT_DIR = ROOT / "analysis" / "out" / "simzfish_fidelity_audit_20260603"


@dataclass(frozen=True)
class ComponentClaim:
    component: str
    original_evidence: str
    current_evidence: str
    status: str
    score: float
    interpretation: str
    required_next_step: str


def _read(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _line_for(path: Path, pattern: str) -> int | None:
    rx = re.compile(pattern)
    for i, line in enumerate(_read(path).splitlines(), start=1):
        if rx.search(line):
            return i
    return None


def _line_ref(path: Path, pattern: str, label: str) -> str:
    line = _line_for(path, pattern)
    rel = str(path.resolve())
    if line is None:
        return f"{label}: {rel}"
    return f"{label}: {rel}:{line}"


def _extract_defines(path: Path) -> dict[str, str]:
    out: dict[str, str] = {}
    for line in _read(path).splitlines():
        m = re.match(r"\s*#define\s+([A-Za-z_][A-Za-z0-9_]*)\s+(.+?)\s*(?://.*)?$", line)
        if m:
            out[m.group(1)] = m.group(2).strip()
    return out


def _load_params() -> dict[str, Any]:
    path = SIMZFISH_RAW / "params.csv"
    rows: list[dict[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            rows.append(row)
    names = [row.get("Parameter name", "") for row in rows if row.get("Parameter name")]
    model_a = []
    model_c = []
    for row in rows:
        try:
            model_a.append(float(row.get("Parameter value - Model A", "nan")))
            model_c.append(float(row.get("Parameter value - Model C", "nan")))
        except ValueError:
            pass
    return {
        "path": str(path.resolve()),
        "parameter_count": len(rows),
        "first_parameters": names[:16],
        "model_a_min": float(np.nanmin(model_a)) if model_a else None,
        "model_a_max": float(np.nanmax(model_a)) if model_a else None,
        "model_c_min": float(np.nanmin(model_c)) if model_c else None,
        "model_c_max": float(np.nanmax(model_c)) if model_c else None,
    }


def _workbook_summary(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    total_nonempty = 0
    total_numeric = 0
    numeric_values: list[float] = []
    for ws in wb.worksheets:
        nonempty = 0
        numeric = 0
        local_values: list[float] = []
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if value is None:
                    continue
                nonempty += 1
                if isinstance(value, (int, float)) and math.isfinite(float(value)):
                    numeric += 1
                    if len(local_values) < 20000:
                        local_values.append(float(value))
        total_nonempty += nonempty
        total_numeric += numeric
        numeric_values.extend(local_values[: max(0, 20000 - len(numeric_values))])
        sheets.append(
            {
                "name": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_cells": nonempty,
                "numeric_cells": numeric,
            }
        )
    wb.close()
    return {
        "path": str(path.relative_to(SIMZFISH_DATA)),
        "size": path.stat().st_size,
        "sheet_count": len(sheets),
        "nonempty_cells": total_nonempty,
        "numeric_cells": total_numeric,
        "numeric_sample_mean": float(np.mean(numeric_values)) if numeric_values else None,
        "numeric_sample_std": float(np.std(numeric_values)) if numeric_values else None,
        "sheets": sheets[:12],
    }


def _mat_summary(path: Path) -> dict[str, Any]:
    summary: dict[str, Any] = {
        "path": str(path.relative_to(SIMZFISH_DATA)),
        "size": path.stat().st_size,
        "loader": "",
        "arrays": [],
        "nested_numeric_cells": [],
    }
    try:
        data = scipy_io.loadmat(path, simplify_cells=True)
        summary["loader"] = "scipy.io.loadmat(simplify_cells=True)"
        for key, value in sorted(data.items()):
            if key.startswith("__"):
                continue
            arr = np.asarray(value)
            summary["arrays"].append(
                {
                    "name": key,
                    "shape": list(arr.shape),
                    "dtype": str(arr.dtype),
                    "numeric": bool(np.issubdtype(arr.dtype, np.number)),
                }
            )
            if arr.dtype == object:
                for idx, item in enumerate(arr.ravel()):
                    item_arr = np.asarray(item)
                    if item_arr.size and np.issubdtype(item_arr.dtype, np.number):
                        finite = item_arr[np.isfinite(item_arr.astype(float, copy=False))]
                        summary["nested_numeric_cells"].append(
                            {
                                "parent": key,
                                "cell_index": int(idx),
                                "shape": list(item_arr.shape),
                                "dtype": str(item_arr.dtype),
                                "min": float(np.nanmin(finite)) if finite.size else None,
                                "max": float(np.nanmax(finite)) if finite.size else None,
                                "mean": float(np.nanmean(finite)) if finite.size else None,
                            }
                        )
    except Exception as exc:  # noqa: BLE001
        summary["loader"] = f"failed: {type(exc).__name__}: {exc}"
    return summary


def _summarize_data_files() -> dict[str, Any]:
    xlsx_files = sorted(SIMZFISH_DATA.rglob("*.xlsx"))
    mat_files = sorted(SIMZFISH_DATA.rglob("*.mat"))
    workbook_summaries = [_workbook_summary(path) for path in xlsx_files]
    mat_summaries = [_mat_summary(path) for path in mat_files]
    mat_nested_numeric_cells = sum(len(item.get("nested_numeric_cells", [])) for item in mat_summaries)
    category_counts: dict[str, dict[str, int]] = {}
    for path in [*xlsx_files, *mat_files]:
        rel = path.relative_to(SIMZFISH_DATA)
        category = rel.parts[0] if len(rel.parts) > 1 else "root"
        entry = category_counts.setdefault(category, {"files": 0, "bytes": 0})
        entry["files"] += 1
        entry["bytes"] += path.stat().st_size
    return {
        "xlsx_count": len(xlsx_files),
        "mat_count": len(mat_files),
        "mat_nested_numeric_cells": mat_nested_numeric_cells,
        "total_files": len(xlsx_files) + len(mat_files),
        "total_bytes": sum(path.stat().st_size for path in [*xlsx_files, *mat_files]),
        "categories": category_counts,
        "workbooks": workbook_summaries,
        "mat_files": mat_summaries,
    }


def _current_config_values() -> dict[str, Any]:
    text = _read(CURRENT_CONFIG)
    values: dict[str, Any] = {}
    for name in [
        "N_BODY_SEGMENTS",
        "BODY_LENGTH_M",
        "BODY_RADIUS_M",
        "TAIL_BEAT_FREQ_MIN_HZ",
        "TAIL_BEAT_FREQ_MAX_HZ",
        "TAIL_BEAT_PHASE_LAG_RAD",
        "TAIL_BEAT_MAX_AMPLITUDE_RAD",
    ]:
        m = re.search(rf"^{name}\s*=\s*(.+)$", text, flags=re.MULTILINE)
        if not m:
            continue
        raw = m.group(1).strip()
        try:
            values[name] = float(raw)
            if values[name].is_integer():
                values[name] = int(values[name])
        except Exception:  # noqa: BLE001
            values[name] = raw
    return values


def _source_summary() -> dict[str, Any]:
    files = {
        "Image.c": SIMZFISH_RAW / "Image.c",
        "OMR.c": SIMZFISH_RAW / "OMR.c",
        "LeakyIntegrator.c": SIMZFISH_RAW / "LeakyIntegrator.c",
        "Robot.c": SIMZFISH_RAW / "Robot.c",
        "ZebrafishRobot_mini.c": SIMZFISH_RAW / "ZebrafishRobot_mini.c",
        "Image.h": SIMZFISH_RAW / "Image.h",
    }
    return {
        "original_defines": {name: _extract_defines(path) for name, path in files.items()},
        "params": _load_params(),
        "line_anchors": {
            "original_two_cameras": _line_ref(files["ZebrafishRobot_mini.c"], r"left_camera, right_camera", "two Webots cameras"),
            "original_retina_update_every_25_cycles": _line_ref(
                files["ZebrafishRobot_mini.c"], r"num_of_cycles % 25", "retina update cadence"
            ),
            "original_off_bipolar": _line_ref(files["Image.c"], r"bipolarCellsComputation", "OFF bipolar cells"),
            "original_dsc": _line_ref(files["Image.c"], r"DSCellsComputation", "direction-selective cells"),
            "original_omr_params": _line_ref(files["OMR.c"], r"readModelParametersFromCSV", "parameterized OMR network"),
            "original_leaky_integrator": _line_ref(files["LeakyIntegrator.c"], r"SS_threshold", "bout leaky integrator"),
            "original_motor_targets": _line_ref(files["Robot.c"], r"getTargetMotorPositions", "motor target generation"),
            "current_video_pipeline": _line_ref(CURRENT_VIDEO, r"calcOpticalFlowFarneback", "backend Farneback flow"),
            "current_simzfish_retina": _line_ref(CURRENT_OMR, r"class SimZFishRetina", "current retina adapter"),
            "current_simzfish_action": _line_ref(CURRENT_OMR, r"action_from_frame", "current OMR action adapter"),
            "current_tail_latent": _line_ref(CURRENT_ACTION, r"tail_targets_from_scores", "current common tail latent"),
        },
    }


def _component_claims(source: dict[str, Any], data: dict[str, Any]) -> list[ComponentClaim]:
    anchors = source["line_anchors"]
    params = source["params"]
    cfg = _current_config_values()
    return [
        ComponentClaim(
            component="visual input geometry",
            original_evidence=(
                f"Webots uses two 320x240 cameras and lower-field sampling; {anchors['original_two_cameras']}; "
                f"Image.h constants {source['original_defines']['Image.h']}"
            ),
            current_evidence=f"Single backend video frame is resized to 160x90 and split into left/right fields; {anchors['current_video_pipeline']}",
            status="partial",
            score=0.45,
            interpretation="The current branch has two-eye features, but not real binocular cameras or calibrated lens geometry.",
            required_next_step="Use the simZFish/ZBot camera geometry or a calibrated larval two-eye projection before claiming visual equivalence.",
        ),
        ComponentClaim(
            component="OFF bipolar retinal transform",
            original_evidence=f"Original uses sigmoid(frame[t-1]-frame[t]-bias 3) over uint8 frames; {anchors['original_off_bipolar']}",
            current_evidence=f"Current adapter uses normalized previous-current OFF sigmoid with scaled bias; {anchors['current_simzfish_retina']}",
            status="near_equivalent",
            score=0.80,
            interpretation="The basic OFF contrast transform is scale-equivalent, though preprocessing and input distribution differ.",
            required_next_step="Regression-test normalized Python OFF output against a direct port of Image.c on synthetic frame pairs.",
        ),
        ComponentClaim(
            component="direction-selective retinal cells",
            original_evidence=(
                "Original Barlow-Levick-style DSC compares OFF_BC_3D current and delayed adjacent pixels with GAP=3/GAPbi=2; "
                f"{anchors['original_dsc']}"
            ),
            current_evidence=f"Current adapter uses residual optical flow after camera-motion subtraction; {anchors['current_video_pipeline']}",
            status="partial",
            score=0.45,
            interpretation="Both estimate motion direction, but the current flow model is not an exact Image.c DSC port.",
            required_next_step="Implement an Image.c-compatible DSC path and compare it with the current flow path across the same videos.",
        ),
        ComponentClaim(
            component="pretectum and hindbrain OMR network",
            original_evidence=(
                f"Original loads {params['parameter_count']} Model A/C parameters from params.csv and computes PT, MLF, and LHB populations; "
                f"{anchors['original_omr_params']}"
            ),
            current_evidence=f"Current action adapter has low-dimensional left/right PT, turn_state, and bout_state approximations; {anchors['current_simzfish_action']}",
            status="missing_exact_port",
            score=0.25,
            interpretation="The current controller is motif-level, not parameter-equivalent to OMR.c.",
            required_next_step="Port OMR.c parameter loading and population equations into Python and make deterministic C-vs-Python unit tests.",
        ),
        ComponentClaim(
            component="bout leaky integrator",
            original_evidence=(
                "Original uses stochastic SS_MLF threshold 436, 500-countdown bouts, command probabilities, and left/right vSPN currents; "
                f"{anchors['original_leaky_integrator']}"
            ),
            current_evidence="Current video adapter emits kick/force/side continuously through a smoothed bout_state.",
            status="partial",
            score=0.35,
            interpretation="The current branch can generate bouts but does not reproduce the original burst-glide integrator.",
            required_next_step="Port LeakyIntegrator.c with seeded RNG and compare bout interval, duration, command, and amplitude distributions.",
        ),
        ComponentClaim(
            component="tail motor target generation",
            original_evidence=(
                "Original computes CPG/vSPN-modulated targets for 7 segment states and actuates 6 tail motors; "
                f"{anchors['original_motor_targets']}"
            ),
            current_evidence=(
                f"Current lab uses {cfg.get('N_BODY_SEGMENTS')} MuJoCo tail segments, {cfg.get('TAIL_BEAT_FREQ_MIN_HZ')}-"
                f"{cfg.get('TAIL_BEAT_FREQ_MAX_HZ')} Hz tail-frequency range, and shared tail_targets_from_scores; "
                f"{anchors['current_tail_latent']}"
            ),
            status="partial",
            score=0.40,
            interpretation="The current MuJoCo body has higher segment resolution but is not calibrated to original simZFish motor targets.",
            required_next_step="Fit or map original 6-motor/7-CPG targets onto the 16-segment MuJoCo body and validate tail-angle distributions.",
        ),
        ComponentClaim(
            component="published calibration data use",
            original_evidence=(
                f"Cached Data_Liu_simZFish_2025 has {data['total_files']} files, "
                f"{data['total_bytes']} bytes, {data['xlsx_count']} workbooks, and {data['mat_count']} MAT files."
            ),
            current_evidence="Current 50k study uses the resources as audit evidence, but does not yet fit behavior parameters from those tables.",
            status="available_not_integrated",
            score=0.30,
            interpretation="The published data are available locally and rich enough for calibration, but not yet wired into parameter fitting.",
            required_next_step="Extract OMR, rheotaxis, lens, ZBot, neural, and tail-behavior targets and optimize current MuJoCo parameters against them.",
        ),
        ComponentClaim(
            component="ZAPBench and simZFish bridge",
            original_evidence="ZAPBench provides whole-brain calcium and fictive tail ephys; Z-Robot/simZFish provides OMR embodiment resources.",
            current_evidence="The lab joins both branches through a shared action latent and MuJoCo body; no paired public dataset links both directly.",
            status="engineering_bridge",
            score=0.50,
            interpretation="The bridge is transparent and useful, but it is not a single experimentally validated chain.",
            required_next_step="Acquire paired visual stimulus, whole-brain calcium, motor ephys, and tail/body kinematics for the same preparation.",
        ),
    ]


def _write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _plot_scorecard(rows: list[ComponentClaim], path: Path) -> None:
    labels = [row.component for row in rows]
    scores = [row.score for row in rows]
    colors = [
        "#23a55a" if row.score >= 0.75 else "#d8a31a" if row.score >= 0.45 else "#c94b4b"
        for row in rows
    ]
    fig, ax = plt.subplots(figsize=(11, 5.8))
    y = np.arange(len(labels))
    ax.barh(y, scores, color=colors)
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlim(0, 1)
    ax.set_xlabel("implementation fidelity score (0=missing, 1=exact/validated)")
    ax.set_title("Current Python/MuJoCo Video Branch vs Public simZFish Controller")
    for i, score in enumerate(scores):
        ax.text(score + 0.02, i, f"{score:.2f}", va="center", fontsize=9)
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def _plot_dataset_inventory(data: dict[str, Any], path: Path) -> None:
    categories = data["categories"]
    names = list(categories)
    counts = [categories[name]["files"] for name in names]
    sizes = [categories[name]["bytes"] / 1e6 for name in names]
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
    axes[0].barh(np.arange(len(names)), counts, color="#377eb8")
    axes[0].set_yticks(np.arange(len(names)), names)
    axes[0].invert_yaxis()
    axes[0].set_title("File Count")
    axes[0].set_xlabel("files")
    axes[1].barh(np.arange(len(names)), sizes, color="#4daf4a")
    axes[1].set_yticks(np.arange(len(names)), names)
    axes[1].invert_yaxis()
    axes[1].set_title("Data Size")
    axes[1].set_xlabel("MB")
    fig.suptitle("Published Data_Liu_simZFish_2025 Inventory")
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def _plot_tail_mapping(path: Path) -> None:
    original = np.array([0.202952981, 0.36415025, 0.481232303, 0.593806825, 0.69596167, 1.0])
    current_segments = 16
    frac = np.linspace(0, 1, current_segments)
    current_envelope = 0.22 + 0.92 * frac
    current_envelope /= np.max(current_envelope)
    fig, ax = plt.subplots(figsize=(9, 4.4))
    ax.plot(np.linspace(0, 1, len(original)), original / np.max(original), marker="o", label="simZFish seg_ampl")
    ax.plot(frac, current_envelope, marker="s", label="current 16-segment envelope")
    ax.set_xlabel("normalized rostral-to-caudal tail coordinate")
    ax.set_ylabel("normalized amplitude")
    ax.set_title("Original simZFish Motor Envelope vs Current MuJoCo Tail Envelope")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def _report(
    *,
    source: dict[str, Any],
    data: dict[str, Any],
    claims: list[ComponentClaim],
    artifacts: dict[str, str],
) -> str:
    mean_score = float(np.mean([c.score for c in claims]))
    status_counts: dict[str, int] = {}
    for claim in claims:
        status_counts[claim.status] = status_counts.get(claim.status, 0) + 1
    rows = "\n".join(
        "| {component} | {status} | {score:.2f} | {interpretation} |".format(
            component=c.component,
            status=c.status,
            score=c.score,
            interpretation=c.interpretation,
        )
        for c in claims
    )
    required = "\n".join(f"{i + 1}. {c.required_next_step}" for i, c in enumerate(claims))
    categories = "\n".join(
        f"- `{name}`: {entry['files']} files, {entry['bytes'] / 1e6:.2f} MB"
        for name, entry in data["categories"].items()
    )
    return f"""# simZFish / Z-Robot Implementation Fidelity Audit

## Working Conclusion

The current zebrafish video branch is **paper-inspired and instrumented**, but
it is **not yet an exact simZFish/Z-Robot port**.  The mean component fidelity
score is `{mean_score:.2f}` on a 0-1 scale.  The strongest overlap is the OFF
retinal contrast transform; the weakest areas are the parameterized OMR network,
the stochastic leaky-integrator bout generator, and direct calibration against
the published Data_Liu_simZFish_2025 tables.

Status counts: `{json.dumps(status_counts, sort_keys=True)}`.

## Visual Scorecard

![Fidelity scorecard]({artifacts['scorecard_png']})

## Published Dataset Inventory

![Dataset inventory]({artifacts['dataset_inventory_png']})

{categories}

The cached public data contain `{data['total_files']}` behavior/neural/resource
files (`{data['total_bytes']}` bytes): `{data['xlsx_count']}` spreadsheets and
`{data['mat_count']}` MAT files.  The MAT files expose
`{data['mat_nested_numeric_cells']}` nested numeric behavior cells after
`scipy.io.loadmat(..., simplify_cells=True)`.

## Tail-Motor Mapping Check

![Tail mapping]({artifacts['tail_mapping_png']})

The current MuJoCo model has higher tail resolution than simZFish, which is
useful for rendering and hydrodynamic control, but it still needs a calibrated
mapping from the original six actuated Webots motors / seven CPG segment states
onto the 16-segment MuJoCo muscle layout.

## Component Evidence Matrix

| component | status | score | interpretation |
|---|---|---:|---|
{rows}

## Original simZFish Anchors

- Two-camera controller: `{source['line_anchors']['original_two_cameras']}`
- Retina update cadence: `{source['line_anchors']['original_retina_update_every_25_cycles']}`
- OFF bipolar cells: `{source['line_anchors']['original_off_bipolar']}`
- Direction-selective cells: `{source['line_anchors']['original_dsc']}`
- Parameterized OMR network: `{source['line_anchors']['original_omr_params']}`
- Leaky integrator: `{source['line_anchors']['original_leaky_integrator']}`
- Motor target generation: `{source['line_anchors']['original_motor_targets']}`
- Params CSV: `{source['params']['path']}` with `{source['params']['parameter_count']}` parameter rows.

## Current Implementation Anchors

- Backend stabilized optical flow: `{source['line_anchors']['current_video_pipeline']}`
- Current retina adapter: `{source['line_anchors']['current_simzfish_retina']}`
- Current OMR action adapter: `{source['line_anchors']['current_simzfish_action']}`
- Common MuJoCo tail latent: `{source['line_anchors']['current_tail_latent']}`

## Required Experiments For Journal-Level Equivalence

{required}

## Generated Artifacts

- Component scorecard CSV: `{artifacts['component_csv']}`
- Original/current source summary JSON: `{artifacts['source_summary_json']}`
- Published data summary JSON: `{artifacts['data_summary_json']}`
- Scorecard figure: `{artifacts['scorecard_png']}`
- Dataset inventory figure: `{artifacts['dataset_inventory_png']}`
- Tail mapping figure: `{artifacts['tail_mapping_png']}`

## Lab-Presentation Boundary

The defensible statement is that the current system is a high-observability
MuJoCo zebrafish lab with a simZFish-inspired video branch.  A stronger claim
that it reproduces the Z-Robot/simZFish controller requires exact porting or
input-output equivalence tests for `Image.c`, `OMR.c`, `LeakyIntegrator.c`, and
`Robot.c`, followed by calibration against the published simZFish/ZBot behavior
tables and DANDI calcium motifs.
"""


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    source = _source_summary()
    data = _summarize_data_files()
    claims = _component_claims(source, data)
    component_rows = [
        {
            "component": c.component,
            "status": c.status,
            "score": f"{c.score:.4f}",
            "original_evidence": c.original_evidence,
            "current_evidence": c.current_evidence,
            "interpretation": c.interpretation,
            "required_next_step": c.required_next_step,
        }
        for c in claims
    ]
    component_csv = OUT_DIR / "component_fidelity_scorecard.csv"
    source_json = OUT_DIR / "source_code_evidence.json"
    data_json = OUT_DIR / "published_data_inventory.json"
    scorecard_png = OUT_DIR / "component_fidelity_scorecard.png"
    dataset_png = OUT_DIR / "published_dataset_inventory.png"
    tail_png = OUT_DIR / "tail_motor_mapping_comparison.png"
    report_path = OUT_DIR / "SIMZFISH_FIDELITY_AUDIT.md"
    _write_csv(
        component_csv,
        component_rows,
        [
            "component",
            "status",
            "score",
            "original_evidence",
            "current_evidence",
            "interpretation",
            "required_next_step",
        ],
    )
    source_json.write_text(json.dumps(source, indent=2, sort_keys=True), encoding="utf-8")
    data_json.write_text(json.dumps(data, indent=2, sort_keys=True), encoding="utf-8")
    _plot_scorecard(claims, scorecard_png)
    _plot_dataset_inventory(data, dataset_png)
    _plot_tail_mapping(tail_png)
    artifacts = {
        "component_csv": str(component_csv.resolve()),
        "source_summary_json": str(source_json.resolve()),
        "data_summary_json": str(data_json.resolve()),
        "scorecard_png": str(scorecard_png.resolve()),
        "dataset_inventory_png": str(dataset_png.resolve()),
        "tail_mapping_png": str(tail_png.resolve()),
    }
    report_path.write_text(
        _report(source=source, data=data, claims=claims, artifacts=artifacts),
        encoding="utf-8",
    )
    manifest = {
        "out_dir": str(OUT_DIR.resolve()),
        "report": str(report_path.resolve()),
        "artifacts": artifacts,
        "mean_score": float(np.mean([c.score for c in claims])),
        "component_count": len(claims),
        "data_file_count": data["total_files"],
        "data_bytes": data["total_bytes"],
    }
    (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    print(json.dumps(manifest, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
