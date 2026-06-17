"""Extract simZFish/Z-Robot calibration targets and compare current 50k runs.

The fidelity audit says which implementation pieces are exact or approximate.
This script turns the published Z-Robot/simZFish data cache into numeric
calibration targets: bout frequencies/fractions, neural response summaries,
rheotaxis traces, ZBot statistics, and MAT behavior-cell summaries.  It then
compares those targets with the current zebrafish MuJoCo 50k-tick runs.
"""

from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy import io as scipy_io


ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = ROOT / "analysis" / "cache" / "z_robot" / "simzfish_data" / "Data_Liu_simZFish_2025"
RUN_ROOT = ROOT / "analysis" / "out" / "comprehensive_activity_study" / "20260603_verified_50k"
OUT_DIR = ROOT / "analysis" / "out" / "simzfish_calibration_targets_20260603"


def _clean(value: Any) -> str:
    return str(value or "").strip().replace("\n", " ")


def _num(value: Any) -> float | None:
    if isinstance(value, (int, float, np.number)):
        out = float(value)
        if np.isfinite(out):
            return out
    return None


def _mean(values: list[float]) -> float:
    return float(np.mean(values)) if values else float("nan")


def _std(values: list[float]) -> float:
    return float(np.std(values)) if values else float("nan")


def _p(values: list[float], q: float) -> float:
    return float(np.percentile(values, q)) if values else float("nan")


def _safe_div(a: float, b: float) -> float:
    return float(a / b) if b else float("nan")


def _find_header(rows: list[tuple[int, list[Any]]], required: tuple[str, ...]) -> tuple[int, list[str]] | None:
    for row_i, values in rows:
        labels = [_clean(v).lower() for v in values]
        text = " ".join(labels)
        if all(term.lower() in text for term in required):
            return row_i, labels
    return None


def _all_rows(ws, max_cols: int | None = None) -> list[tuple[int, list[Any]]]:
    rows: list[tuple[int, list[Any]]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = list(row[:max_cols]) if max_cols is not None else list(row)
        rows.append((i, values))
    return rows


def _idx(labels: list[str], *terms: str) -> int | None:
    for i, label in enumerate(labels):
        if all(term.lower() in label for term in terms):
            return i
    return None


def _prob_idx(labels: list[str], direction: str) -> int | None:
    for i, label in enumerate(labels):
        lower = label.lower()
        if direction.lower() in lower and ("probability" in lower or "possibility" in lower):
            return i
    return None


def _dataset_group(path: Path) -> str:
    rel = path.relative_to(DATA_ROOT)
    name = path.name.lower()
    if rel.parts[0] == "Lens_effects":
        return "lens_effects"
    if rel.parts[0] == "Rheotaxis_simZFish":
        return "rheotaxis"
    if rel.parts[0] == "Connectivity_effects":
        return "connectivity_effects"
    if rel.parts[0] == "Figure6_ZBotStats":
        return "zbot_stats"
    if rel.parts[0] == "Zebrafish_ParallelStimOMRBehavior":
        return "zebrafish_behavior_mat"
    if "free swimming omr" in name or "free swimming" in name:
        return "free_swimming_omr"
    if "neuronal recording" in name or "neuronal" in name:
        return "neuronal_recording"
    return "root"


@dataclass
class LocomotionSheet:
    workbook: str
    sheet: str
    group: str
    rows: int
    duration_s: float
    bout_frequency_hz: float
    left_bouts: int
    forward_bouts: int
    right_bouts: int
    left_fraction: float
    forward_fraction: float
    right_fraction: float
    prob_forward_mean: float
    prob_left_mean: float
    prob_right_mean: float
    lnmlf_mean: float
    rnmlf_mean: float
    lahb_mean: float
    rahb_mean: float
    x_displacement: float
    y_displacement: float
    xy_path: float


def _extract_locomotion_sheet(path: Path, ws) -> LocomotionSheet | None:
    rows = _all_rows(ws, max_cols=24)
    header = _find_header(rows[:8], ("time", "action"))
    if header is None:
        return None
    header_row, labels = header
    time_i = _idx(labels, "time")
    action_i = _idx(labels, "action")
    pf_i = _prob_idx(labels, "forward")
    pl_i = _prob_idx(labels, "left")
    pr_i = _prob_idx(labels, "right")
    ln_i = _idx(labels, "lnmlf")
    rn_i = _idx(labels, "rnmlf")
    la_i = _idx(labels, "lahb")
    ra_i = _idx(labels, "rahb")
    x_i = _idx(labels, "position x")
    y_i = _idx(labels, "position y")
    if time_i is None or action_i is None:
        return None

    times: list[float] = []
    actions: list[int] = []
    pf: list[float] = []
    pl: list[float] = []
    pr: list[float] = []
    ln: list[float] = []
    rn: list[float] = []
    la: list[float] = []
    ra: list[float] = []
    xs: list[float] = []
    ys: list[float] = []
    for row_i, values in rows:
        if row_i <= header_row:
            continue
        t = _num(values[time_i] if time_i < len(values) else None)
        action = _num(values[action_i] if action_i < len(values) else None)
        if t is None or action is None:
            continue
        times.append(t)
        actions.append(int(round(action)))
        for idx_value, target in [(pf_i, pf), (pl_i, pl), (pr_i, pr), (ln_i, ln), (rn_i, rn), (la_i, la), (ra_i, ra)]:
            if idx_value is not None and idx_value < len(values):
                value = _num(values[idx_value])
                if value is not None:
                    target.append(value)
        if x_i is not None and x_i < len(values):
            x = _num(values[x_i])
            if x is not None:
                xs.append(x)
        if y_i is not None and y_i < len(values):
            y = _num(values[y_i])
            if y is not None:
                ys.append(y)
    if not times:
        return None
    duration = max(times) - min(times)
    counts = Counter(actions)
    total = len(actions)
    path_len = float("nan")
    if len(xs) >= 2 and len(ys) >= 2:
        dx = np.diff(np.asarray(xs, dtype=float))
        dy = np.diff(np.asarray(ys, dtype=float))
        path_len = float(np.sum(np.sqrt(dx * dx + dy * dy)))
    return LocomotionSheet(
        workbook=str(path.relative_to(DATA_ROOT)),
        sheet=ws.title,
        group=_dataset_group(path),
        rows=total,
        duration_s=duration,
        bout_frequency_hz=_safe_div(total, duration),
        left_bouts=int(counts[-1]),
        forward_bouts=int(counts[0]),
        right_bouts=int(counts[1]),
        left_fraction=_safe_div(counts[-1], total),
        forward_fraction=_safe_div(counts[0], total),
        right_fraction=_safe_div(counts[1], total),
        prob_forward_mean=_mean(pf),
        prob_left_mean=_mean(pl),
        prob_right_mean=_mean(pr),
        lnmlf_mean=_mean(ln),
        rnmlf_mean=_mean(rn),
        lahb_mean=_mean(la),
        rahb_mean=_mean(ra),
        x_displacement=(xs[-1] - xs[0]) if len(xs) >= 2 else float("nan"),
        y_displacement=(ys[-1] - ys[0]) if len(ys) >= 2 else float("nan"),
        xy_path=path_len,
    )


def _extract_neural_sheet(path: Path, ws) -> list[dict[str, Any]]:
    rows = _all_rows(ws, max_cols=64)
    header = _find_header(rows[:8], ("time",))
    if header is None:
        return []
    header_row, labels = header
    time_i = _idx(labels, "time")
    if time_i is None:
        return []
    data_by_col: dict[int, list[float]] = defaultdict(list)
    for row_i, values in rows:
        if row_i <= header_row:
            continue
        t = _num(values[time_i] if time_i < len(values) else None)
        if t is None:
            continue
        for col_i, label in enumerate(labels):
            if col_i == time_i:
                continue
            if not label:
                continue
            value = _num(values[col_i] if col_i < len(values) else None)
            if value is not None:
                data_by_col[col_i].append(value)
    out: list[dict[str, Any]] = []
    for col_i, values in data_by_col.items():
        label = labels[col_i] or f"col_{col_i}"
        if len(values) < 10:
            continue
        out.append(
            {
                "workbook": str(path.relative_to(DATA_ROOT)),
                "sheet": ws.title,
                "group": _dataset_group(path),
                "channel": label,
                "column": col_i,
                "n": len(values),
                "mean": _mean(values),
                "std": _std(values),
                "p95": _p(values, 95),
                "max": float(np.max(values)),
            }
        )
    return out


def _extract_grid_trace_sheet(path: Path, ws) -> list[dict[str, Any]]:
    """Extract rheotaxis-like traces with a Time column and several trial/condition columns."""
    rows = _all_rows(ws, max_cols=32)
    header: tuple[int, list[str]] | None = None
    for row_i, values in rows[:8]:
        labels = [_clean(v) for v in values]
        if any("time" in label.lower() for label in labels) and sum(v is not None for v in values) >= 3:
            header = (row_i, labels)
            break
    if header is None:
        return []
    header_row, labels_raw = header
    labels = [str(label) for label in labels_raw]
    time_i = next((i for i, label in enumerate(labels) if "time" in label.lower()), None)
    if time_i is None:
        return []
    by_col: dict[int, list[tuple[float, float]]] = defaultdict(list)
    for row_i, values in rows:
        if row_i <= header_row:
            continue
        t = _num(values[time_i] if time_i < len(values) else None)
        if t is None:
            continue
        for col_i, label in enumerate(labels):
            if col_i == time_i or not _clean(label):
                continue
            value = _num(values[col_i] if col_i < len(values) else None)
            if value is not None:
                by_col[col_i].append((t, value))
    out: list[dict[str, Any]] = []
    for col_i, pairs in by_col.items():
        if len(pairs) < 20:
            continue
        ts = np.asarray([p[0] for p in pairs], dtype=float)
        vs = np.asarray([p[1] for p in pairs], dtype=float)
        duration = float(ts[-1] - ts[0])
        slope = float((vs[-1] - vs[0]) / duration) if duration else float("nan")
        out.append(
            {
                "workbook": str(path.relative_to(DATA_ROOT)),
                "sheet": ws.title,
                "group": _dataset_group(path),
                "condition": labels[col_i],
                "n": len(pairs),
                "duration_s": duration,
                "start": float(vs[0]),
                "end": float(vs[-1]),
                "delta": float(vs[-1] - vs[0]),
                "mean": float(np.mean(vs)),
                "std": float(np.std(vs)),
                "slope_per_s": slope,
            }
        )
    return out


def _extract_zbot_stats(path: Path, ws) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for _, values in _all_rows(ws, max_cols=32):
        labels = [_clean(v) for v in values]
        if "No motor" in labels and "No OMR" in labels and "OMR" in labels:
            header = labels
            continue
        if "header" not in locals():
            continue
        nums = [_num(v) for v in values]
        if sum(v is not None for v in nums) >= 2:
            for label, value in zip(header, nums):
                if label and value is not None:
                    out.append(
                        {
                            "workbook": str(path.relative_to(DATA_ROOT)),
                            "sheet": ws.title,
                            "condition": label,
                            "value": value,
                        }
                    )
    return out


def _extract_workbooks() -> dict[str, list[dict[str, Any]]]:
    locomotion: list[dict[str, Any]] = []
    neural: list[dict[str, Any]] = []
    grid: list[dict[str, Any]] = []
    zbot: list[dict[str, Any]] = []
    for path in sorted(DATA_ROOT.rglob("*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                loco = _extract_locomotion_sheet(path, ws)
                if loco is not None:
                    locomotion.append(loco.__dict__)
                # Neural sheets share the same simple time + many channel layout.
                title = (path.name + " " + ws.title).lower()
                if "neuronal" in title or "dsgc" in title or "recording" in title:
                    neural.extend(_extract_neural_sheet(path, ws))
                if _dataset_group(path) == "rheotaxis":
                    grid.extend(_extract_grid_trace_sheet(path, ws))
                if _dataset_group(path) == "zbot_stats":
                    zbot.extend(_extract_zbot_stats(path, ws))
        finally:
            wb.close()
    return {
        "locomotion": locomotion,
        "neural": neural,
        "rheotaxis_grid": grid,
        "zbot_stats": zbot,
    }


def _extract_mats() -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for path in sorted((DATA_ROOT / "Zebrafish_ParallelStimOMRBehavior").glob("*.mat")):
        data = scipy_io.loadmat(path, simplify_cells=True)
        behave = np.asarray(data.get("BeHave", []), dtype=object)
        numeric_cells = []
        total_numeric_elements = 0
        for idx, value in enumerate(behave.ravel()):
            arr = np.asarray(value)
            if arr.size and np.issubdtype(arr.dtype, np.number):
                values = arr.astype(float, copy=False)
                finite = values[np.isfinite(values)]
                if finite.size:
                    total_numeric_elements += int(finite.size)
                    numeric_cells.append(
                        {
                            "cell_index": idx,
                            "shape": list(arr.shape),
                            "n": int(finite.size),
                            "mean": float(np.mean(finite)),
                            "std": float(np.std(finite)),
                            "min": float(np.min(finite)),
                            "max": float(np.max(finite)),
                        }
                    )
        out.append(
            {
                "file": str(path.relative_to(DATA_ROOT)),
                "cells": int(behave.size),
                "numeric_cell_count": len(numeric_cells),
                "numeric_elements": total_numeric_elements,
                "numeric_cells": numeric_cells,
            }
        )
    return out


def _read_csv_dicts(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def _read_actions(path: Path) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                out.append(json.loads(line))
    return out


def _current_run_summary(run: str) -> dict[str, Any]:
    run_dir = RUN_ROOT / run
    state_rows = _read_csv_dicts(run_dir / "state_metrics.csv")
    actions = _read_actions(run_dir / "action_samples.jsonl")
    ticks = [float(row["tick"]) for row in state_rows if row.get("tick")]
    sim_times = [float(row["sim_time_s"]) for row in state_rows if row.get("sim_time_s")]
    metrics = {
        "run": run,
        "frames": len(state_rows),
        "tick_span": int(max(ticks) - min(ticks)) if ticks else 0,
        "duration_s": float(max(sim_times) - min(sim_times)) if sim_times else 0.0,
        "actions": len(actions),
    }
    for field in ["speed_xy_mm_s", "heading_rate_rad_s", "body_max_local_bend_2d_rad", "body_straightness", "tail_yaw_rms", "muscle_abs_max"]:
        vals = [float(row[field]) for row in state_rows if row.get(field) not in (None, "")]
        metrics[f"{field}_mean"] = _mean(vals)
        metrics[f"{field}_p95"] = _p(vals, 95)
    metrics["xy_displacement_mm"] = float(state_rows[-1]["xy_displacement_mm"]) if state_rows else float("nan")
    action_types = Counter()
    noncoast_indices: list[int] = []
    noncoast_forces: list[float] = []
    for i, action in enumerate(actions):
        state = action.get("state", action)
        bout_type = state.get("action_bout_type") or state.get("bout_type") or state.get("side") or "none"
        kick = float(state.get("action_kick", state.get("kick", 0.0)) or 0.0)
        force = float(state.get("action_force", state.get("force", 0.0)) or 0.0)
        action_types[str(bout_type)] += 1
        if str(bout_type) not in {"none", "coast"} and (kick >= 0.5 or force > 0.05):
            noncoast_indices.append(i)
            noncoast_forces.append(force)
    events = 0
    prev = -999999
    for idx in noncoast_indices:
        if idx - prev > 1:
            events += 1
        prev = idx
    duration = metrics["duration_s"] or 1.0
    metrics["noncoast_frame_fraction"] = _safe_div(len(noncoast_indices), len(actions))
    metrics["noncoast_event_count"] = events
    metrics["noncoast_event_frequency_hz"] = _safe_div(events, duration)
    metrics["noncoast_command_frequency_hz"] = _safe_div(len(noncoast_indices), duration)
    metrics["action_force_mean_noncoast"] = _mean(noncoast_forces)
    metrics["action_type_counts"] = dict(action_types)
    total_noncoast = max(1, sum(count for key, count in action_types.items() if key not in {"none", "coast"}))
    metrics["left_like_fraction"] = _safe_div(action_types.get("omr_turn_left", 0) + action_types.get("left", 0), total_noncoast)
    metrics["right_like_fraction"] = _safe_div(action_types.get("omr_turn_right", 0) + action_types.get("right", 0), total_noncoast)
    metrics["forward_like_fraction"] = _safe_div(
        action_types.get("omr_forward_bout", 0) + action_types.get("fictive_tail_bout", 0), total_noncoast
    )
    metrics["startle_like_fraction"] = _safe_div(action_types.get("startle_c_bend", 0), total_noncoast)
    return metrics


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fields: list[str] = []
    seen = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                fields.append(key)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _plot_bout_fractions(locomotion: list[dict[str, Any]], current: list[dict[str, Any]], path: Path) -> None:
    selected = [row for row in locomotion if row["group"] in {"free_swimming_omr", "lens_effects"}]
    by_group: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in selected:
        by_group[row["group"]].append(row)
    rows: list[tuple[str, float, float, float]] = []
    for group, group_rows in by_group.items():
        rows.append(
            (
                group,
                _mean([r["left_fraction"] for r in group_rows]),
                _mean([r["forward_fraction"] for r in group_rows]),
                _mean([r["right_fraction"] for r in group_rows]),
            )
        )
    for row in current:
        rows.append(
            (
                "current_" + row["run"],
                row.get("left_like_fraction", 0.0),
                row.get("forward_like_fraction", 0.0),
                row.get("right_like_fraction", 0.0),
            )
        )
    labels = [r[0] for r in rows]
    left = np.asarray([r[1] for r in rows], dtype=float)
    forward = np.asarray([r[2] for r in rows], dtype=float)
    right = np.asarray([r[3] for r in rows], dtype=float)
    y = np.arange(len(rows))
    fig, ax = plt.subplots(figsize=(11, 5.5))
    ax.barh(y, left, label="left", color="#4e79a7")
    ax.barh(y, forward, left=left, label="forward", color="#59a14f")
    ax.barh(y, right, left=left + forward, label="right", color="#e15759")
    ax.set_yticks(y, labels)
    ax.set_xlim(0, 1)
    ax.set_xlabel("fraction")
    ax.set_title("Published simZFish Bout Fractions vs Current Action Labels")
    ax.legend(loc="lower right")
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def _plot_frequency(locomotion: list[dict[str, Any]], current: list[dict[str, Any]], path: Path) -> None:
    groups = ["free_swimming_omr", "lens_effects", "root"]
    data = [[row["bout_frequency_hz"] for row in locomotion if row["group"] == group] for group in groups]
    labels = groups[:]
    for row in current:
        data.append([row.get("noncoast_event_frequency_hz", float("nan"))])
        labels.append("current_event_" + row["run"])
        data.append([row.get("noncoast_command_frequency_hz", float("nan"))])
        labels.append("current_command_" + row["run"])
    fig, ax = plt.subplots(figsize=(12, 5.5))
    ax.boxplot(data, tick_labels=labels, showfliers=False)
    ax.set_ylabel("Hz")
    ax.set_title("Published Bout Frequency vs Current Event/Command Frequency")
    ax.tick_params(axis="x", rotation=35)
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def _plot_neural(neural: list[dict[str, Any]], path: Path) -> None:
    top = sorted(neural, key=lambda r: float(r["mean"]), reverse=True)[:40]
    labels = [f"{Path(r['workbook']).stem[:18]}:{r['sheet'][:12]}:{r['channel']}" for r in top]
    means = [float(r["mean"]) for r in top]
    p95 = [float(r["p95"]) for r in top]
    y = np.arange(len(top))
    fig, ax = plt.subplots(figsize=(12, 10))
    ax.barh(y, p95, color="#f28e2b", label="p95")
    ax.barh(y, means, color="#76b7b2", label="mean")
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlabel("activation")
    ax.set_title("Top Published simZFish Neural Activation Targets")
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def _plot_rheotaxis(grid: list[dict[str, Any]], path: Path) -> None:
    selected = [r for r in grid if r["group"] == "rheotaxis"]
    selected = sorted(selected, key=lambda r: (r["workbook"], r["sheet"], str(r["condition"])))[:80]
    labels = [f"{Path(r['workbook']).stem[:16]}:{r['sheet']}:{r['condition']}" for r in selected]
    deltas = [float(r["delta"]) for r in selected]
    y = np.arange(len(selected))
    fig, ax = plt.subplots(figsize=(12, max(4, len(selected) * 0.16)))
    ax.barh(y, deltas, color="#9c755f")
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlabel("end-start value")
    ax.set_title("Published Rheotaxis/Heading Trace Deltas")
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def _comparison_rows(locomotion: list[dict[str, Any]], current: list[dict[str, Any]]) -> list[dict[str, Any]]:
    published_freqs = [float(row["bout_frequency_hz"]) for row in locomotion if row["group"] in {"free_swimming_omr", "lens_effects", "root"}]
    published_left = [float(row["left_fraction"]) for row in locomotion if row["group"] in {"free_swimming_omr", "lens_effects", "root"}]
    published_forward = [float(row["forward_fraction"]) for row in locomotion if row["group"] in {"free_swimming_omr", "lens_effects", "root"}]
    published_right = [float(row["right_fraction"]) for row in locomotion if row["group"] in {"free_swimming_omr", "lens_effects", "root"}]
    rows: list[dict[str, Any]] = []
    for run in current:
        event_freq = float(run.get("noncoast_event_frequency_hz", float("nan")))
        command_freq = float(run.get("noncoast_command_frequency_hz", float("nan")))
        rows.append(
            {
                "run": run["run"],
                "target_family": "published_simzfish_locomotion",
                "published_bout_frequency_mean_hz": _mean(published_freqs),
                "published_bout_frequency_p05_hz": _p(published_freqs, 5),
                "published_bout_frequency_p95_hz": _p(published_freqs, 95),
                "current_noncoast_event_frequency_hz": event_freq,
                "current_noncoast_command_frequency_hz": command_freq,
                "event_frequency_ratio_to_published_mean": _safe_div(event_freq, _mean(published_freqs)),
                "command_frequency_ratio_to_published_mean": _safe_div(command_freq, _mean(published_freqs)),
                "published_left_fraction_mean": _mean(published_left),
                "published_forward_fraction_mean": _mean(published_forward),
                "published_right_fraction_mean": _mean(published_right),
                "current_left_like_fraction": run.get("left_like_fraction", float("nan")),
                "current_forward_like_fraction": run.get("forward_like_fraction", float("nan")),
                "current_right_like_fraction": run.get("right_like_fraction", float("nan")),
                "current_startle_like_fraction": run.get("startle_like_fraction", float("nan")),
                "interpretation": "event frequency is transition-based; command frequency is frame/row occupancy and should not be treated as bout rate",
            }
        )
    return rows


def _report(
    workbook_data: dict[str, list[dict[str, Any]]],
    mat_rows: list[dict[str, Any]],
    current: list[dict[str, Any]],
    comparisons: list[dict[str, Any]],
    artifacts: dict[str, str],
) -> str:
    locomotion = workbook_data["locomotion"]
    neural = workbook_data["neural"]
    grid = workbook_data["rheotaxis_grid"]
    zbot = workbook_data["zbot_stats"]
    published_freqs = [r["bout_frequency_hz"] for r in locomotion if r["group"] in {"free_swimming_omr", "lens_effects", "root"}]
    comp_lines = "\n".join(
        f"- `{row['run']}`: event Hz `{row['current_noncoast_event_frequency_hz']:.4f}`, command Hz `{row['current_noncoast_command_frequency_hz']:.4f}`, published mean Hz `{row['published_bout_frequency_mean_hz']:.4f}`."
        for row in comparisons
    )
    current_lines = "\n".join(
        f"- `{row['run']}`: {row['tick_span']} ticks, {row['duration_s']:.3f}s, noncoast events `{row['noncoast_event_count']}`, action types `{json.dumps(row['action_type_counts'], sort_keys=True)}`."
        for row in current
    )
    return f"""# simZFish / Z-Robot Calibration Target Atlas

## Working Conclusion

The published Z-Robot/simZFish resources contain enough numeric targets to
start calibrating the current MuJoCo video branch.  This pass extracted
locomotion, neural, rheotaxis, ZBot, and MAT behavior summaries from the cached
public data and compared them with the current 50k zebrafish runs.

The most important finding is that current video action occupancy is not a
published bout-rate match by itself.  The fairer transition-based event rate
must be used for calibration, while frame/row command occupancy is an internal
controller diagnostic.

## Extracted Target Coverage

- Locomotion/bout sheets: `{len(locomotion)}`
- Neural channel summaries: `{len(neural)}`
- Rheotaxis/grid trace summaries: `{len(grid)}`
- ZBot scalar stats: `{len(zbot)}`
- MAT behavior files: `{len(mat_rows)}`
- MAT numeric behavior cells: `{sum(row['numeric_cell_count'] for row in mat_rows)}`
- Published locomotion bout frequency mean: `{_mean(published_freqs):.4f}` Hz, p05-p95 `{_p(published_freqs, 5):.4f}`-`{_p(published_freqs, 95):.4f}` Hz

## Current Run Comparison

{current_lines}

{comp_lines}

## Visual Calibration Views

![Bout fractions]({artifacts['bout_fractions_png']})

![Bout frequency]({artifacts['bout_frequency_png']})

![Neural activation targets]({artifacts['neural_targets_png']})

![Rheotaxis deltas]({artifacts['rheotaxis_png']})

## Interpretation For The Lab

The current data package can now support a real calibration plan:

1. Use locomotion sheets to target bout-rate and left/forward/right command fractions.
2. Use neural sheets to target PT/MLF/LHB-like activation motifs.
3. Use rheotaxis sheets to target displacement and heading responses under flow.
4. Use ZBot tables as physical robot sanity targets.
5. Use MAT behavior cells for detailed bout/heading/trajectory statistics after field-label reverse engineering.

The current implementation still needs a fitting loop.  This atlas establishes
the target surface, not a completed calibration.

## Generated Tables

- Locomotion targets: `{artifacts['locomotion_csv']}`
- Neural targets: `{artifacts['neural_csv']}`
- Rheotaxis/grid targets: `{artifacts['rheotaxis_csv']}`
- ZBot stats: `{artifacts['zbot_csv']}`
- MAT summary: `{artifacts['mat_json']}`
- Current run metrics: `{artifacts['current_csv']}`
- Current-vs-published comparison: `{artifacts['comparison_csv']}`
"""


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook_data = _extract_workbooks()
    mat_rows = _extract_mats()
    current = [_current_run_summary("calcium_all"), _current_run_summary("video_commons_tenggol_underwater")]
    comparisons = _comparison_rows(workbook_data["locomotion"], current)

    locomotion_csv = OUT_DIR / "published_locomotion_targets.csv"
    neural_csv = OUT_DIR / "published_neural_targets.csv"
    rheotaxis_csv = OUT_DIR / "published_rheotaxis_targets.csv"
    zbot_csv = OUT_DIR / "published_zbot_stats.csv"
    current_csv = OUT_DIR / "current_50k_run_metrics.csv"
    comparison_csv = OUT_DIR / "current_vs_published_behavior.csv"
    mat_json = OUT_DIR / "published_mat_behavior_summary.json"
    _write_csv(locomotion_csv, workbook_data["locomotion"])
    _write_csv(neural_csv, workbook_data["neural"])
    _write_csv(rheotaxis_csv, workbook_data["rheotaxis_grid"])
    _write_csv(zbot_csv, workbook_data["zbot_stats"])
    _write_csv(current_csv, current)
    _write_csv(comparison_csv, comparisons)
    mat_json.write_text(json.dumps(mat_rows, indent=2, sort_keys=True), encoding="utf-8")

    bout_fractions_png = OUT_DIR / "published_vs_current_bout_fractions.png"
    bout_frequency_png = OUT_DIR / "published_vs_current_bout_frequency.png"
    neural_png = OUT_DIR / "published_neural_activation_targets.png"
    rheotaxis_png = OUT_DIR / "published_rheotaxis_trace_deltas.png"
    _plot_bout_fractions(workbook_data["locomotion"], current, bout_fractions_png)
    _plot_frequency(workbook_data["locomotion"], current, bout_frequency_png)
    _plot_neural(workbook_data["neural"], neural_png)
    _plot_rheotaxis(workbook_data["rheotaxis_grid"], rheotaxis_png)

    artifacts = {
        "locomotion_csv": str(locomotion_csv.resolve()),
        "neural_csv": str(neural_csv.resolve()),
        "rheotaxis_csv": str(rheotaxis_csv.resolve()),
        "zbot_csv": str(zbot_csv.resolve()),
        "current_csv": str(current_csv.resolve()),
        "comparison_csv": str(comparison_csv.resolve()),
        "mat_json": str(mat_json.resolve()),
        "bout_fractions_png": str(bout_fractions_png.resolve()),
        "bout_frequency_png": str(bout_frequency_png.resolve()),
        "neural_targets_png": str(neural_png.resolve()),
        "rheotaxis_png": str(rheotaxis_png.resolve()),
    }
    report_path = OUT_DIR / "SIMZFISH_CALIBRATION_TARGET_ATLAS.md"
    report_path.write_text(
        _report(
            workbook_data=workbook_data,
            mat_rows=mat_rows,
            current=current,
            comparisons=comparisons,
            artifacts=artifacts,
        ),
        encoding="utf-8",
    )
    manifest = {
        "out_dir": str(OUT_DIR.resolve()),
        "report": str(report_path.resolve()),
        "artifacts": artifacts,
        "locomotion_targets": len(workbook_data["locomotion"]),
        "neural_targets": len(workbook_data["neural"]),
        "rheotaxis_targets": len(workbook_data["rheotaxis_grid"]),
        "zbot_stats": len(workbook_data["zbot_stats"]),
        "mat_files": len(mat_rows),
        "mat_numeric_cells": sum(row["numeric_cell_count"] for row in mat_rows),
        "comparisons": len(comparisons),
    }
    (OUT_DIR / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    print(json.dumps(manifest, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
